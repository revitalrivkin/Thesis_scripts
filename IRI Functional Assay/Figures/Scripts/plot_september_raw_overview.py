"""
4.1.3 IRI Functional Assay Results - September Experiment, Raw Viability Overview
================================================================================
Raw-data overview for the September 2026 IRI experiment, built to match the
July raw-overview figure (plot_july_raw_overview.py) as closely as possible:
same house style (bar = mean, error bar = SD, jittered per-well dots), same
palette roles (Baseline=green, IRI ctrl=grey, Healthy EEV/NegEV=blue shades,
Diabetic EEV/NegEV=red/orange shades), same "hard exclusions first, then
generic Tukey's-fence per column" approach.

September has a wider donor set (6 Healthy, 6 Diabetic columns before
exclusions) and several STRUCTURAL exclusions beyond simple statistical
outliers - all taken from the annotated ground truth
`Plasma EVs\IRI\September IRI\data\September_IRI_plate_layout.xlsx` and
note 11.2's "Finalized Analysis" section, confirmed 2026-09-21:

  1. Diabetic columns 1-4 (3T EEV, 3T NegEV, P4 EEV, P4 NegEV) - EXCLUDED
     ENTIRELY. Pipetting error (150uL CellTiter-Glo instead of 100uL) on
     cols 1-3; col 4 (P4 NegEV) excluded for matched-pair integrity once
     its EEV partner (col 3) is gone. Net effect: 3T and P4 do not appear
     in this figure at all - usable Diabetic donors are P5, P6, P7, 1T.
  2. Diabetic column 12 (IRI control), row A (0.169M) - excluded as a
     pipetting/edge error; NOT caught by Tukey's fence on its own (verified:
     the fence on the remaining 7 values is [0.14, 13.94]M, which would not
     exclude 0.169M if it were still in the set - this exclusion is
     structural, not statistical, and must be hard-coded). IRI control mean
     uses rows B-H only (n=7, mean 7.1957M - reproduces note 11.2's stated
     value exactly).
  3. Healthy plate's own IRI control (column 12) - LEFT OUT OF THIS FIGURE
     ENTIRELY per Revital's explicit call (2026-09-21), even though the raw
     values exist - it's not used as a reference anywhere (IRI had no
     detectable effect there, ~19.8M = baseline) and showing it would imply
     it's part of the analysis when it isn't.
  4. Healthy column 3 row H (H10 EEV) and column 11 row H (H1 EEV) - the
     annotated file calls these "reallocated to No-treatment control pool,"
     but note 11.2's finalized IRI control mean is explicitly n=7 (diabetic
     only) and does NOT include them - the two numbers don't reconcile.
     Per Revital's explicit call (2026-09-21, option "c"), both wells are
     simply DROPPED rather than guessing where they belong: H10 EEV and H1
     EEV are each analyzed with 7 wells (row H excluded), not reallocated
     anywhere.

All other exclusions are handled by the same generic Tukey's-fence filter
used in July (verified 2026-09-21 to exactly reproduce every individual
outlier flagged in the annotated ground-truth file, including the baseline
exclusion - no need to hand-list them).

Data source: ../../Data/September/06.09.26 Viabilty Healthy plate.xlsx,
             ../../Data/September/06.09.26 Viabilty Diabetic plate.xlsx,
             ../../Data/September/06.09.26 Viabilty No-IRI contro.l skax.xlsx

Output: September_Raw_Overview.png/.svg (one level up, in Figures/)
"""

import os
import numpy as np
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

FONT = "Arial"
BG = "#fcfcfb"
GRID = "#e1e0d9"
SPINE = "#c3c2b7"
TITLE_COLOR = "#0b0b0b"
LABEL_COLOR = "#52514e"

HERE = os.path.dirname(os.path.abspath(__file__))
FIG_DIR = os.path.dirname(HERE)
DATA_DIR = os.path.join(FIG_DIR, "..", "Data", "September")
XLSX_HEALTHY = os.path.join(DATA_DIR, "06.09.26 Viabilty Healthy plate.xlsx")
XLSX_DIABETIC = os.path.join(DATA_DIR, "06.09.26 Viabilty Diabetic plate.xlsx")
XLSX_BASELINE = os.path.join(DATA_DIR, "06.09.26 Viabilty No-IRI contro.l skax.xlsx")

RNG_SEED = 42
ROWS = list("ABCDEFGH")


def read_plate(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Luminescence 1_01"]
    raw = {}
    for ri, r in enumerate(ROWS):
        for col in range(1, 13):
            raw[(r, col)] = ws.cell(row=11 + ri, column=col + 1).value
    return raw


def tukey_filter(vals):
    a = np.asarray(vals, dtype=float)
    if len(a) < 4:
        return a
    q1, q3 = np.percentile(a, 25), np.percentile(a, 75)
    iqr = q3 - q1
    lo, hi = q1 - 1.5 * iqr, q3 + 1.5 * iqr
    return a[(a >= lo) & (a <= hi)]


healthy_raw = read_plate(XLSX_HEALTHY)
diabetic_raw = read_plate(XLSX_DIABETIC)


def get_wells_M(raw, col, drop_rows=()):
    vals = []
    for r in ROWS:
        if r in drop_rows:
            continue
        v = raw.get((r, col))
        if v is not None:
            vals.append(v / 1e6)
    return np.array(vals)


# ── Baseline: row A only, columns 1-6 (cols 7-12 are background noise, not
# baseline wells - confirmed from raw file: near-zero/negative values) ──────
wb_base = openpyxl.load_workbook(XLSX_BASELINE, data_only=True)
ws_base = wb_base["Luminescence 1_01"]
baseline_raw = [ws_base.cell(row=11, column=c + 1).value / 1e6 for c in range(1, 7)]
baseline_vals = tukey_filter(baseline_raw)

# ── IRI control: diabetic plate col 12, row A hard-excluded (pipetting/edge
# error, not caught by Tukey), then Tukey filter on rows B-H (no-op, verified) ─
iri_ctrl_raw = get_wells_M(diabetic_raw, 12, drop_rows={"A"})
iri_ctrl_vals = tukey_filter(iri_ctrl_raw)

# ── Group definitions, in display order (matches July's structure) ─────────
C_BASE, C_BASE_DOT = "#3ba272", "#c9ead9"
C_CTRL, C_CTRL_DOT = "#7a8694", "#dadde1"
C_H_EEV, C_H_EEV_DOT = "#2a78d6", "#cde2fb"
C_H_NEG, C_H_NEG_DOT = "#6aadff", "#d9ebff"
C_D_EEV, C_D_EEV_DOT = "#e34948", "#f5b8b8"
C_D_NEG, C_D_NEG_DOT = "#eb6834", "#fbd2bb"

groups = [
    ("Baseline\n(no IRI)", baseline_vals, C_BASE, C_BASE_DOT),
    ("IRI ctrl\n(no EVs)", iri_ctrl_vals, C_CTRL, C_CTRL_DOT),
    # Healthy - column order per plate layout: H11, H10, H2, H3, H4, H1
    ("H11\nEEV", tukey_filter(get_wells_M(healthy_raw, 1)), C_H_EEV, C_H_EEV_DOT),
    ("H11\nNegEV", tukey_filter(get_wells_M(healthy_raw, 2)), C_H_NEG, C_H_NEG_DOT),
    ("H10\nEEV", tukey_filter(get_wells_M(healthy_raw, 3, drop_rows={"H"})), C_H_EEV, C_H_EEV_DOT),
    ("H10\nNegEV", tukey_filter(get_wells_M(healthy_raw, 4)), C_H_NEG, C_H_NEG_DOT),
    ("H2\nEEV", tukey_filter(get_wells_M(healthy_raw, 5)), C_H_EEV, C_H_EEV_DOT),
    ("H2\nNegEV", tukey_filter(get_wells_M(healthy_raw, 6)), C_H_NEG, C_H_NEG_DOT),
    ("H3\nEEV", tukey_filter(get_wells_M(healthy_raw, 7)), C_H_EEV, C_H_EEV_DOT),
    ("H3\nNegEV", tukey_filter(get_wells_M(healthy_raw, 8)), C_H_NEG, C_H_NEG_DOT),
    ("H4\nEEV", tukey_filter(get_wells_M(healthy_raw, 9)), C_H_EEV, C_H_EEV_DOT),
    ("H4\nNegEV", tukey_filter(get_wells_M(healthy_raw, 10)), C_H_NEG, C_H_NEG_DOT),
    ("H1\nEEV", tukey_filter(get_wells_M(healthy_raw, 11, drop_rows={"H"})), C_H_EEV, C_H_EEV_DOT),
    # Diabetic - columns 1-4 (3T, P4) excluded entirely; usable: P5, P6, P7, 1T
    # Display labels use the unified D-naming convention (1T->D1, P5->D5,
    # P6->D6, P7->D7) - Healthy stays H.
    ("D5\nEEV", tukey_filter(get_wells_M(diabetic_raw, 5)), C_D_EEV, C_D_EEV_DOT),
    ("D5\nNegEV", tukey_filter(get_wells_M(diabetic_raw, 6)), C_D_NEG, C_D_NEG_DOT),
    ("D6\nEEV", tukey_filter(get_wells_M(diabetic_raw, 7)), C_D_EEV, C_D_EEV_DOT),
    ("D6\nNegEV", tukey_filter(get_wells_M(diabetic_raw, 8)), C_D_NEG, C_D_NEG_DOT),
    ("D7\nEEV", tukey_filter(get_wells_M(diabetic_raw, 9)), C_D_EEV, C_D_EEV_DOT),
    ("D7\nNegEV", tukey_filter(get_wells_M(diabetic_raw, 10)), C_D_NEG, C_D_NEG_DOT),
    ("D1\nEEV", tukey_filter(get_wells_M(diabetic_raw, 11)), C_D_EEV, C_D_EEV_DOT),
]

rng = np.random.default_rng(RNG_SEED)
fig, ax = plt.subplots(figsize=(19, 6.2))
fig.patch.set_facecolor(BG)
ax.set_facecolor(BG)

x = np.arange(len(groups))
means = [g[1].mean() for g in groups]
sds = [g[1].std(ddof=1) for g in groups]
bar_colors = [g[2] for g in groups]
dot_colors = [g[3] for g in groups]
ns = [len(g[1]) for g in groups]

ax.bar(x, means, yerr=sds,
       color=bar_colors, edgecolor=bar_colors, linewidth=0.8,
       error_kw=dict(elinewidth=1.4, capsize=4, capthick=1.4, ecolor=LABEL_COLOR),
       width=0.6, zorder=3, alpha=0.80)

for xi, (label, vals, _, dcolor) in zip(x, groups):
    jit = rng.uniform(-0.14, 0.14, len(vals))
    ax.scatter(np.full(len(vals), xi) + jit, vals, color=dcolor,
               edgecolors="white", linewidths=0.5, s=36, zorder=5)

BASELINE_MEAN = baseline_vals.mean()
IRI_CTRL_MEAN = iri_ctrl_vals.mean()
ax.axhline(BASELINE_MEAN, color=C_BASE, linestyle=(0, (4, 3)), linewidth=1.1, zorder=2)
ax.axhline(IRI_CTRL_MEAN, color=C_CTRL, linestyle=(0, (4, 3)), linewidth=1.1, zorder=2)

ax.set_axisbelow(True)
ax.yaxis.grid(True, color=GRID, linewidth=0.8, zorder=0)
ax.set_xticks(x)
ax.set_xticklabels([f"{g[0]}\n(n={n})" for g, n in zip(groups, ns)],
                    fontsize=15, fontfamily=FONT, color=TITLE_COLOR)
ax.set_xlim(-0.6, len(groups) - 0.4)
ax.set_ylabel("Raw Luminescence (millions)", fontsize=18, fontfamily=FONT,
              color=LABEL_COLOR, labelpad=10)
ax.tick_params(axis="y", labelsize=14, labelcolor=LABEL_COLOR)
ax.tick_params(axis="x", bottom=False)
for lbl in ax.get_yticklabels():
    lbl.set_fontfamily(FONT)
for spine in ["top", "right", "bottom"]:
    ax.spines[spine].set_visible(False)
ax.spines["left"].set_color(SPINE)
ax.spines["left"].set_linewidth(0.8)

trans = ax.get_xaxis_transform()
ax.text(7.5, -0.27, "Healthy donors", ha="center", va="top", fontsize=18,
        fontweight="bold", color=C_H_EEV, fontfamily=FONT, transform=trans, clip_on=False)
ax.text(16.5, -0.27, "Diabetic donors", ha="center", va="top", fontsize=18,
        fontweight="bold", color=C_D_EEV, fontfamily=FONT, transform=trans, clip_on=False)

ax.set_title("September 2026 IRI Experiment - Raw Viability Signal, All Groups",
             fontsize=23, fontweight="bold", fontfamily=FONT,
             color=TITLE_COLOR, pad=18, loc="center")

fig.tight_layout()
for ext in ("png", "svg"):
    fig.savefig(os.path.join(FIG_DIR, f"September_Raw_Overview.{ext}"), dpi=200,
                bbox_inches="tight", facecolor=fig.get_facecolor())
plt.close(fig)

print(f"Baseline mean: {BASELINE_MEAN:.3f} M (n={len(baseline_vals)}, dropped {len(baseline_raw)-len(baseline_vals)} outlier)")
print(f"IRI ctrl mean: {IRI_CTRL_MEAN:.3f} M (n={len(iri_ctrl_vals)})")
print("\nGroup n's:")
for g, n in zip(groups, ns):
    print(f"  {g[0].replace(chr(10),' '):<16} n={n}")
print("\nSaved: September_Raw_Overview.png/.svg")
